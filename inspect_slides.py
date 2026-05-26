from pptx import Presentation

prs = Presentation(r'C:\Users\loeff\Documents\EPFL\semester-project\LoRa_Presentation_Technique.pptx')
slides = list(prs.slides)
print(f'Total slides: {len(slides)}')

for i, slide in enumerate(slides):
    print(f'\n=== Slide {i+1} ===')
    print(f'  Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()[:80]
            if text:
                print(f'  [{shape.name}] {repr(text)}')
