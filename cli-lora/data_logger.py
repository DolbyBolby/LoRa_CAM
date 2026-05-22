"""
data_logger.py — Logger pour le sweep de paramètres LoRa SX1280.

Protocole série reçu depuis le firmware receive_blocking :
  COMP_START|{comp_id}|{sf}|{bw}|{cr}
  PKT_RX|{comp_id}|{seq}|{rssi}|{snr}|{freq_err}
  SWEEP_DONE

Génère un fichier Excel avec 3 feuilles :
  - Raw Data   : un rang par slot de paquet attendu (reçu ou perdu)
  - Summary    : une ligne par composition avec statistiques
  - Heatmaps   : matrices RSSI par BW (SF × CR)
"""

import serial
import csv
import re
import time
from datetime import datetime
from pathlib import Path
from statistics import mean, stdev

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

# ---- Dossier de sortie par défaut ----
DATA_DIR = Path(__file__).parent / "data"

# ---- Constantes sweep ----
PACKETS_PER_COMP    = 50           # Paquets par composition (50 → PER à 2% de résolution)
TOTAL_COMPOSITIONS  = 8 * 4 * 4   # SF5-12 × BW×4 × CR×4 = 128

# ---- Couleurs (ARGB) ----
_HDR_BG    = "FF1F4E79"
_COMP_BG   = "FFD6E4F0"
_LOST_BG   = "FFFFCCCC"
_RED_PER   = "FFFF9999"
_ORA_PER   = "FFFFCC99"
_GRN_PER   = "FF99FF99"


def _resolve_path(output_file: str) -> Path:
    p = Path(output_file)
    # Les chemins relatifs (peu importe les sous-dossiers passés) → DATA_DIR/nom_fichier
    # Les chemins absolus sont utilisés tels quels
    if not p.is_absolute():
        p = DATA_DIR / p.name
    p.parent.mkdir(parents=True, exist_ok=True)
    return p.resolve()  # Toujours absolu


# ===========================================================================
# Simple capture (ancienne fonctionnalité — gardée pour la commande log-data)
# ===========================================================================

_RSSI_RE = re.compile(r"^RSSI:\s+(-?\d+\.?\d*)")
_SNR_RE  = re.compile(r"^SNR:\s+(-?\d+\.?\d*)")
_FREQ_RE = re.compile(r"^Frequency Error:\s+(-?\d+\.?\d*)")
_COLUMNS_SIMPLE = ["#", "RSSI (dBm)", "SNR (dB)", "Frequency Error (Hz)"]


def capture_and_save_data(
    num_packets: int,
    port: str = "COM5",
    baudrate: int = 9600,
    output_file: str = "data_log.csv",
):
    """Capture N paquets (RSSI/SNR/FreqErr) depuis le firmware RX simple."""
    records, current, packet_count = [], {}, 0
    output_path = _resolve_path(output_file)

    print(f"Écoute sur {port} à {baudrate} baud...")
    print(f"Capture de {num_packets} paquets → {output_path}\n")

    try:
        with serial.Serial(port=port, baudrate=baudrate, timeout=2) as ser:
            while packet_count < num_packets:
                raw = ser.readline()
                if not raw:
                    continue
                line = raw.decode("utf-8", errors="ignore").strip()
                if not line:
                    continue

                m = _RSSI_RE.match(line)
                if m:
                    current["RSSI (dBm)"] = float(m.group(1))
                    continue
                m = _SNR_RE.match(line)
                if m:
                    current["SNR (dB)"] = float(m.group(1))
                    continue
                m = _FREQ_RE.match(line)
                if m:
                    current["Frequency Error (Hz)"] = float(m.group(1))
                    if len(current) == 3:
                        packet_count += 1
                        current["#"] = packet_count
                        records.append(current)
                        print(
                            f"  [{packet_count:>4}/{num_packets}]  "
                            f"RSSI={current['RSSI (dBm)']:>8.2f} dBm  "
                            f"SNR={current['SNR (dB)']:>6.2f} dB  "
                            f"FE={current['Frequency Error (Hz)']:>10.2f} Hz"
                        )
                        current = {}
                    else:
                        current = {}

    except KeyboardInterrupt:
        print("\nCapture arrêtée.")
    except serial.SerialException as e:
        print(f"\nErreur série: {e}")
        return

    if not records:
        print("Aucune donnée capturée.")
        return

    ext = output_path.suffix.lower()
    if ext in (".xlsx", ".xls") and EXCEL_AVAILABLE:
        _simple_save_excel(records, output_path)
    else:
        _simple_save_csv(records, output_path)


def _simple_save_csv(records, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=_COLUMNS_SIMPLE)
        writer.writeheader()
        writer.writerows(records)
    print(f"\nCSV : {path}  ({len(records)} paquets)")


def _simple_save_excel(records, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LoRa Data"
    fill = PatternFill("solid", fgColor=_HDR_BG)
    font = Font(bold=True, color="FFFFFF")
    for ci, col in enumerate(_COLUMNS_SIMPLE, 1):
        c = ws.cell(row=1, column=ci, value=col)
        c.fill = fill; c.font = font
        c.alignment = Alignment(horizontal="center")
    for rec in records:
        ws.append([rec[c] for c in _COLUMNS_SIMPLE])
    wb.save(str(path))
    print(f"\nExcel : {path}  ({len(records)} paquets)")


# ===========================================================================
# Sweep logger
# ===========================================================================

class SweepLogger:
    """Reçoit les lignes série du firmware sweep et génère un fichier Excel."""

    _RAW_COLS = [
        "Composition_ID", "SF", "BW", "CR",
        "Seq_Num", "RSSI", "SNR", "Freq_Error", "Timestamp", "Lost_Packet",
    ]
    _SUM_COLS = [
        "Composition_ID", "SF", "BW", "CR",
        "Packets_Received", "PER(%)",
        "RSSI_Mean", "RSSI_Std", "RSSI_Min", "RSSI_Max",
        "SNR_Mean", "SNR_Std", "Link_Margin_dB",
    ]

    def __init__(self, output_path: Path):
        self.output_path = output_path
        self.sweep_start = time.time()

        self.compositions: list[dict] = []
        self.current_comp: dict | None = None

        # Workbook Excel
        self.wb = openpyxl.Workbook()
        self.ws_raw = self.wb.active
        self.ws_raw.title = "Raw Data"
        self.ws_sum = self.wb.create_sheet("Summary")
        self.ws_hm  = self.wb.create_sheet("Heatmaps")

        self._raw_row = 2
        self._sum_row = 2

        self._write_sheet_headers()

    # ------------------------------------------------------------------
    # Initialisation des en-têtes
    # ------------------------------------------------------------------

    def _hdr_cell(self, ws, row, col, value):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = PatternFill("solid", fgColor=_HDR_BG)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")
        return c

    def _write_sheet_headers(self):
        for ci, col in enumerate(self._RAW_COLS, 1):
            self._hdr_cell(self.ws_raw, 1, ci, col)
        self.ws_raw.freeze_panes = "A2"
        raw_widths = [16, 6, 8, 8, 10, 10, 10, 14, 22, 14]
        for i, w in enumerate(raw_widths, 1):
            self.ws_raw.column_dimensions[get_column_letter(i)].width = w

        for ci, col in enumerate(self._SUM_COLS, 1):
            self._hdr_cell(self.ws_sum, 1, ci, col)
        self.ws_sum.freeze_panes = "A2"
        sum_widths = [16, 6, 8, 8, 18, 10, 12, 12, 12, 12, 10, 10, 16]
        for i, w in enumerate(sum_widths, 1):
            self.ws_sum.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Réception des événements depuis le port série
    # ------------------------------------------------------------------

    def on_comp_start(self, comp_id: int, sf: str, bw: str, cr: str):
        elapsed = int(time.time() - self.sweep_start)
        pct = comp_id / TOTAL_COMPOSITIONS * 100

        # Estimation du temps restant basée sur la vitesse réelle
        if comp_id > 1 and elapsed > 0:
            avg = elapsed / (comp_id - 1)
            remaining = int(avg * (TOTAL_COMPOSITIONS - comp_id + 1))
            eta = f"{remaining // 60}m{remaining % 60:02d}s"
        else:
            eta = "calcul en cours..."

        print(f"\n{'─' * 62}")
        print(f"  COMP {comp_id:>3}/{TOTAL_COMPOSITIONS}  ({pct:5.1f}%)  "
              f"SF={sf}  BW={bw} kHz  CR={cr}")
        print(f"  Écoulé : {elapsed // 60}m{elapsed % 60:02d}s  |  "
              f"ETA restant : {eta}")
        print(f"{'─' * 62}")

        self.current_comp = {
            "id": comp_id, "sf": sf, "bw": bw, "cr": cr,
            "packets": {},   # seq → {rssi, snr, freq_err, ts}
            "finalized": False,
        }
        self.compositions.append(self.current_comp)

    def on_pkt_rx(self, comp_id: int, seq: int,
                  rssi: float, snr: float, freq_err: float):
        if self.current_comp is None or self.current_comp["id"] != comp_id:
            return
        ts = datetime.now().isoformat(timespec="seconds")
        self.current_comp["packets"][seq] = {
            "rssi": rssi, "snr": snr, "freq_err": freq_err, "ts": ts
        }
        received = len(self.current_comp["packets"])
        bar_done = int(received / PACKETS_PER_COMP * 20)
        bar = "█" * bar_done + "░" * (20 - bar_done)
        print(f"  [{received:>2}/{PACKETS_PER_COMP}] {bar}  "
              f"seq={seq:>2}  RSSI={rssi:>7.2f} dBm  "
              f"SNR={snr:>5.2f} dB  FE={freq_err:>9.2f} Hz")

    # ------------------------------------------------------------------
    # Finalisation d'une composition → écriture dans Excel
    # ------------------------------------------------------------------

    def finalize_composition(self, comp: dict):
        if comp.get("finalized"):
            return
        comp["finalized"] = True

        comp_id = comp["id"]
        sf, bw, cr = comp["sf"], comp["bw"], comp["cr"]
        pkts = comp["packets"]  # seq → data

        ws = self.ws_raw

        # Ligne-titre de la composition (fusionnée sur toutes les colonnes)
        ws.merge_cells(
            start_row=self._raw_row, start_column=1,
            end_row=self._raw_row, end_column=len(self._RAW_COLS),
        )
        c = ws.cell(row=self._raw_row, column=1,
                    value=f"=== Composition {comp_id} : SF={sf} BW={bw} CR={cr} ===")
        c.font = Font(bold=True, color="FF1F4E79")
        c.fill = PatternFill("solid", fgColor=_COMP_BG)
        c.alignment = Alignment(horizontal="center")
        self._raw_row += 1

        rssi_vals, snr_vals = [], []
        lost = 0

        for seq in range(PACKETS_PER_COMP):
            if seq in pkts:
                p = pkts[seq]
                rssi_vals.append(p["rssi"])
                snr_vals.append(p["snr"])
                ws.append([comp_id, sf, bw, cr, seq,
                           round(p["rssi"], 2), round(p["snr"], 2),
                           round(p["freq_err"], 2), p["ts"], False])
            else:
                lost += 1
                ws.append([comp_id, sf, bw, cr, seq,
                           None, None, None, None, True])
                for col in range(1, len(self._RAW_COLS) + 1):
                    ws.cell(row=self._raw_row, column=col).fill = \
                        PatternFill("solid", fgColor=_LOST_BG)
            self._raw_row += 1

        self._raw_row += 1  # Ligne vide séparatrice

        # ---- Feuille Summary ----
        received  = PACKETS_PER_COMP - lost
        per_pct   = lost / PACKETS_PER_COMP * 100

        r_mean = round(mean(rssi_vals), 2) if rssi_vals else None
        r_std  = round(stdev(rssi_vals), 2) if len(rssi_vals) > 1 else 0.0
        r_min  = round(min(rssi_vals), 2)  if rssi_vals else None
        r_max  = round(max(rssi_vals), 2)  if rssi_vals else None
        s_mean = round(mean(snr_vals), 2)  if snr_vals else None
        s_std  = round(stdev(snr_vals), 2) if len(snr_vals) > 1 else 0.0
        # Link margin = RSSI - noise floor (on considère -120 dBm comme référence)
        lm = round(r_mean - (-120), 2) if r_mean is not None else None

        self.ws_sum.append([
            comp_id, sf, bw, cr, received, round(per_pct, 2),
            r_mean, r_std, r_min, r_max,
            s_mean, s_std, lm,
        ])

        # Mise en couleur de la cellule PER(%)
        per_cell = self.ws_sum.cell(row=self._sum_row, column=6)
        if per_pct > 10:
            per_cell.fill = PatternFill("solid", fgColor=_RED_PER)
        elif per_pct > 1:
            per_cell.fill = PatternFill("solid", fgColor=_ORA_PER)
        else:
            per_cell.fill = PatternFill("solid", fgColor=_GRN_PER)

        self._sum_row += 1

        # Sauvegarde incrémentale après chaque composition
        self.save()

    # ------------------------------------------------------------------
    # Feuille Heatmaps
    # ------------------------------------------------------------------

    def build_heatmaps(self):
        BW_VALS = ["203", "406", "812", "1625"]
        SF_VALS = ["5", "6", "7", "8", "9", "10", "11", "12"]
        CR_VALS = ["4/5", "4/6", "4/7", "4/8"]
        ws = self.ws_hm

        # Table de lookup (sf, bw, cr) → RSSI moyen
        rssi_map: dict[tuple, float | None] = {}
        for comp in self.compositions:
            vals = [p["rssi"] for p in comp["packets"].values()]
            key = (comp["sf"], comp["bw"], comp["cr"])
            rssi_map[key] = round(mean(vals), 2) if vals else None

        start_col = 1
        for bw in BW_VALS:
            n_cr = len(CR_VALS)
            # Titre BW (fusionné)
            ws.merge_cells(
                start_row=1, start_column=start_col,
                end_row=1, end_column=start_col + n_cr,
            )
            tc = ws.cell(row=1, column=start_col,
                         value=f"BW = {bw} kHz  —  RSSI moyen (dBm)")
            tc.font = Font(bold=True, color="FFFFFF")
            tc.fill = PatternFill("solid", fgColor=_HDR_BG)
            tc.alignment = Alignment(horizontal="center")

            # En-tête CR
            ws.cell(row=2, column=start_col, value="SF \\ CR").font = Font(bold=True)
            for ci, cr in enumerate(CR_VALS):
                c = ws.cell(row=2, column=start_col + 1 + ci, value=cr)
                c.font = Font(bold=True)
                c.alignment = Alignment(horizontal="center")

            # Données SF × CR
            for si, sf in enumerate(SF_VALS):
                ws.cell(row=3 + si, column=start_col,
                        value=f"SF{sf}").font = Font(bold=True)
                for ci, cr in enumerate(CR_VALS):
                    val = rssi_map.get((sf, bw, cr))
                    c = ws.cell(row=3 + si, column=start_col + 1 + ci, value=val)
                    c.alignment = Alignment(horizontal="center")

            # Gradient de couleur : rouge=pire RSSI, vert=meilleur RSSI
            col_start = get_column_letter(start_col + 1)
            col_end   = get_column_letter(start_col + n_cr)
            data_range = f"{col_start}3:{col_end}{2 + len(SF_VALS)}"
            ws.conditional_formatting.add(
                data_range,
                ColorScaleRule(
                    start_type="min",  start_color="FF0000",
                    mid_type="percentile", mid_value=50, mid_color="FFFF00",
                    end_type="max",    end_color="00B050",
                ),
            )

            start_col += n_cr + 2  # Colonne vide entre les matrices

        # Ajuster les largeurs de la feuille Heatmaps
        # col[0] peut être un MergedCell (pas de .column_letter) → on passe par .column
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=6)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max_len + 4

    # ------------------------------------------------------------------
    # Sauvegarde
    # ------------------------------------------------------------------

    def save(self):
        try:
            self.wb.save(str(self.output_path))
        except PermissionError:
            print(f"\n  [ERREUR] Impossible d'écrire : {self.output_path}")
            print("  Le fichier est peut-être ouvert dans Excel — fermez-le et relancez.")
            raise
        except Exception as e:
            print(f"\n  [ERREUR] Sauvegarde échouée : {e}")
            raise

    def print_summary(self):
        print("\n" + "=" * 70)
        print(f"{'ID':>4} {'SF':>4} {'BW':>6} {'CR':>5} "
              f"{'Reçus':>6} {'PER%':>6} {'RSSI_moy':>10} {'SNR_moy':>9}")
        print("-" * 70)
        for comp in self.compositions:
            vals_r = [p["rssi"] for p in comp["packets"].values()]
            vals_s = [p["snr"]  for p in comp["packets"].values()]
            recv = len(comp["packets"])
            per  = (PACKETS_PER_COMP - recv) / PACKETS_PER_COMP * 100
            rm   = f"{mean(vals_r):.1f}" if vals_r else "N/A"
            sm   = f"{mean(vals_s):.1f}" if vals_s else "N/A"
            print(f"{comp['id']:>4} {comp['sf']:>4} {comp['bw']:>6} "
                  f"{comp['cr']:>5} {recv:>6} {per:>6.1f} {rm:>10} {sm:>9}")
        print("=" * 70)


def _fmt_duration(seconds: int) -> str:
    return f"{seconds // 3600}h {(seconds % 3600) // 60}m{seconds % 60:02d}s" \
        if seconds >= 3600 else f"{seconds // 60}m{seconds % 60:02d}s"


def sweep_capture_and_save_data(
    port: str = "COM4",
    baudrate: int = 9600,
    output_file: str = "sweep_results.xlsx",
    verbose: bool = False,
):
    """
    Écoute le port série du firmware receive_blocking en mode sweep
    et génère le fichier Excel complet.
    """
    if not EXCEL_AVAILABLE:
        print("ERREUR : openpyxl requis. Installez avec : pip install openpyxl")
        return

    output_path = _resolve_path(output_file)
    logger = SweepLogger(output_path)

    # Estimation durée : borne basse (sans temps d'émission radio)
    est_low = TOTAL_COMPOSITIONS * (1 + PACKETS_PER_COMP * 0.1 + 1)  # secondes
    print("=" * 62)
    print(f"  SWEEP LoRa SX1280 — {TOTAL_COMPOSITIONS} compositions")
    print(f"  {PACKETS_PER_COMP} paquets/composition")
    print(f"  Durée estimée : >= {_fmt_duration(int(est_low))} "
          f"(plus longue pour SF11/SF12)")
    print(f"  Port : {port} à {baudrate} baud")
    print(f"  Sortie : {output_path}")
    print(f"  Ctrl+C pour sauvegarder et quitter prématurément")
    print("=" * 62 + "\n")

    # Timeout court (1s) pour détecter les périodes d'inactivité
    IDLE_PRINT_INTERVAL = 2.0   # secondes entre deux messages "en attente"
    last_activity   = time.time()
    last_idle_print = 0.0

    def _handle_idle():
        nonlocal last_idle_print
        now = time.time()
        idle = now - last_activity
        if idle > 1.0 and now - last_idle_print > IDLE_PRINT_INTERVAL:
            elapsed = int(now - logger.sweep_start)
            comp_done = len(logger.compositions)
            print(
                f"  ⏳ En attente... {idle:.0f}s  "
                f"[{comp_done}/{TOTAL_COMPOSITIONS} compos terminées | "
                f"écoulé : {_fmt_duration(elapsed)}]   ",
                end="\r",
            )
            last_idle_print = now

    def _clear_idle_line():
        print(" " * 80, end="\r")

    try:
        with serial.Serial(port=port, baudrate=baudrate, timeout=1) as ser:
            while True:
                raw = ser.readline()

                if not raw:
                    _handle_idle()
                    continue

                line = raw.decode("utf-8", errors="ignore").strip()
                if not line:
                    _handle_idle()
                    continue

                _clear_idle_line()
                last_activity = time.time()

                if verbose:
                    print(f"[RAW] {line}")

                # ---- COMP_START|{id}|{sf}|{bw}|{cr} ----
                if line.startswith("COMP_START|"):
                    parts = line.split("|")
                    if len(parts) >= 5:
                        if logger.current_comp:
                            logger.finalize_composition(logger.current_comp)
                        logger.on_comp_start(
                            int(parts[1]), parts[2], parts[3], parts[4]
                        )

                # ---- PKT_RX|{comp_id}|{seq}|{rssi}|{snr}|{freq_err} ----
                elif line.startswith("PKT_RX|"):
                    parts = line.split("|")
                    if len(parts) >= 6:
                        try:
                            logger.on_pkt_rx(
                                int(parts[1]),
                                int(parts[2]),
                                float(parts[3]),
                                float(parts[4]),
                                float(parts[5]),
                            )
                        except ValueError:
                            pass

                # ---- SWEEP_DONE ----
                elif line == "SWEEP_DONE":
                    total_elapsed = int(time.time() - logger.sweep_start)
                    print(f"\n{'=' * 62}")
                    print(f"  SWEEP TERMINÉ en {_fmt_duration(total_elapsed)}")
                    print(f"{'=' * 62}")
                    print("  Finalisation du fichier Excel...")
                    if logger.current_comp:
                        logger.finalize_composition(logger.current_comp)
                    logger.build_heatmaps()
                    logger.save()
                    logger.print_summary()
                    print(f"\n  Fichier sauvegardé : {output_path}")
                    break

    except KeyboardInterrupt:
        total_elapsed = int(time.time() - logger.sweep_start)
        print(f"\n\n  Interrompu après {_fmt_duration(total_elapsed)} "
              f"({len(logger.compositions)}/{TOTAL_COMPOSITIONS} compositions)")
        print("  Sauvegarde des données partielles...")
        if logger.current_comp:
            logger.finalize_composition(logger.current_comp)
        logger.build_heatmaps()
        logger.save()
        logger.print_summary()
        print(f"  Fichier partiel sauvegardé : {output_path}")

    except serial.SerialException as e:
        print(f"\n  Erreur série : {e}")
        if logger.compositions:
            logger.save()
