from pptx import Presentation

prs = Presentation(r'C:\Users\loeff\Documents\EPFL\semester-project\LoRa_Presentation_Technique_v2.pptx')
slides = list(prs.slides)

print("=== New slides (4-11) detail ===")
for i in range(3, 11):
    slide = slides[i]
    print(f"\n--- Slide {i+1} ---")
    n_text = n_pic = n_tbl = n_other = 0
    for shape in slide.shapes:
        try:
            t = shape.shape_type
            if t == 13:
                n_pic += 1
            elif t == 19:
                n_tbl += 1
            elif shape.has_text_frame:
                n_text += 1
            else:
                n_other += 1
        except Exception:
            n_other += 1
    print(f"  texts={n_text}, pictures={n_pic}, tables={n_tbl}, other={n_other}")
    count = 0
    for shape in slide.shapes:
        try:
            if shape.has_text_frame and count < 4:
                txt = shape.text_frame.text.strip()[:70].replace('\n',' ')
                if txt:
                    print(f"  ✓ {txt}")
                    count += 1
        except Exception:
            pass

print("\n=== Slide 1 unchanged check ===")
for shape in slides[0].shapes:
    try:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()[:60]
            if t:
                print(f"  {t}")
    except Exception:
        pass
