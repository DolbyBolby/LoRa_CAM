from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation(r'C:\Users\loeff\Documents\EPFL\semester-project\LoRa_Presentation_Technique.pptx')
slides = list(prs.slides)

def emu_to_in(emu):
    return round(emu / 914400, 4)

def get_fill_color(shape):
    try:
        fill = shape.fill
        if fill.type and fill.type.name == 'SOLID':
            return str(fill.fore_color.rgb)
        return f"type={fill.type}"
    except:
        return "N/A"

# Inspect slides 2 and 3 for exact positions
for slide_idx in [1, 2]:
    slide = slides[slide_idx]
    print(f"\n=== Slide {slide_idx+1} shape positions ===")
    for shape in slide.shapes:
        l = emu_to_in(shape.left) if shape.left else 'N/A'
        t = emu_to_in(shape.top) if shape.top else 'N/A'
        w = emu_to_in(shape.width) if shape.width else 'N/A'
        h = emu_to_in(shape.height) if shape.height else 'N/A'
        fill_col = get_fill_color(shape)
        text_preview = ''
        if shape.has_text_frame:
            text_preview = shape.text_frame.text[:30].replace('\n',' ')
        print(f"  {shape.name}: L={l} T={t} W={w} H={h} | fill={fill_col} | text='{text_preview}'")
