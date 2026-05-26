from pptx import Presentation

prs = Presentation(r'C:\Users\loeff\Documents\EPFL\semester-project\LoRa_Presentation_Technique_v2.pptx')
slides = list(prs.slides)
print(f'Total slides: {len(slides)}')

for i, slide in enumerate(slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()[:60]
            if t:
                texts.append(t)
    first_text = texts[0] if texts else '(no text)'
    print(f'  Slide {i+1:2d}: {first_text}')
