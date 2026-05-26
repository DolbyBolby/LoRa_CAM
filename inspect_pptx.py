from pptx import Presentation
from pptx.util import Pt

prs = Presentation(r'C:\Users\loeff\Documents\EPFL\semester-project\LoRa_Presentation_Technique.pptx')
print(f'Slide count: {len(prs.slides)}')
print(f'Slide width: {prs.slide_width.inches:.2f} in, height: {prs.slide_height.inches:.2f} in')

slides = list(prs.slides)
for i, slide in enumerate(slides[:3]):
    print(f'\n--- Slide {i+1} ---')
    print(f'  Layout: {slide.slide_layout.name}')
    bg = slide.background
    fill = bg.fill
    print(f'  BG fill type: {fill.type}')
    try:
        print(f'  BG color: {fill.fore_color.rgb}')
    except Exception as e:
        print(f'  BG color error: {e}')
    for shape in slide.shapes:
        print(f'  Shape: type={shape.shape_type}, name={shape.name}, has_tf={shape.has_text_frame}')
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    txt = run.text[:50] if run.text else ''
                    font = run.font
                    col = 'N/A'
                    try:
                        if font.color and font.color.type is not None:
                            col = str(font.color.rgb)
                    except:
                        col = 'inherited'
                    print(f'    Run: "{txt}" | bold={font.bold} | size={font.size} | color={col} | name={font.name}')

# Check master background
master = prs.slide_master
bg = master.background
fill = bg.fill
print(f'\n--- Master BG fill type: {fill.type}')
try:
    print(f'  Master BG color: {fill.fore_color.rgb}')
except Exception as e:
    print(f'  Master BG error: {e}')

# Try to get theme element from XML
import lxml.etree as etree
ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
clr_scheme = master.element.find(f'.//{{{ns}}}clrScheme')
if clr_scheme is not None:
    print('\n--- Theme Color Scheme ---')
    for child in clr_scheme:
        tag = child.tag.split('}')[-1]
        for subchild in child:
            subtag = subchild.tag.split('}')[-1]
            val = subchild.get('val') or subchild.get('lastClr') or ''
            print(f'  {tag}: {subtag} = {val}')
